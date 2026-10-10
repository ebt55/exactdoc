"""Markup Word and LibreOffice both render alike (WP21).

Measured in Word 16.0.20430 (desktop, Office 2024) against the canonical
LibreOffice on the same DOCX, 2026-10-05:

- Table edges. Word 2010 layout (compatibilityMode 14) hangs a row left of
  w:tblInd by its first cell's left margin; LibreOffice by the table's default
  left margin. A table written with a zero default and per-cell pads therefore
  stood a pad further left in Word (c1's stat cards 49.6pt). The first column's
  pad is now the table default and part of the indent; a border that hangs left
  of its text column in the source (c3) is kept, in the standard profile.
- Page-number restarts that repeat the previous page's parity under different
  odd/even headers make Word insert a blank page (y19, y25).
- A BaseFont named in a CJK locale's legacy encoding (y51's Shift-JIS
  "ＭＳ ゴシック") reaches the family table.
- The template's Courier and MS Mincho declarations are not on a stock
  Windows + Office machine; the standard profile declares neither.

    python -m unittest tests.test_word_markup
"""
import io
import os
import sys
import tempfile
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import docx  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from exactdoc import docxout, fonts  # noqa: E402
from exactdoc.layout import Cell, HFSection, Para, Run, TableEl  # noqa: E402


def _para(text, left=0.0):
    return Para(runs=[Run(text=text, font="Helvetica", size=10, color="#000000")],
                leading=12.0, left_indent=left)


def _table(pads, left_indent=0.0, hang=0.0, blocks_row=None):
    rows = []
    for i, p in enumerate(pads):
        c0 = Cell(paras=[_para("row %d" % i, left=p)], pad=(2.0, p, 2.0, 4.0))
        if blocks_row == i:
            c0.blocks = list(c0.paras)
        c1 = Cell(paras=[_para("value")], pad=(2.0, 4.0, 2.0, 4.0))
        rows.append([c0, c1])
    return TableEl(rows=rows, col_widths=[120.0, 120.0], row_heights=[None] * len(pads),
                   left_indent=left_indent, hang_left=hang)


def _write(t, profile="standard"):
    d = docx.Document()
    ctx = docxout.WriteCtx(output_profile=profile,
                           line_mode=docxout.line_mode_for(profile))
    tbl = docxout.write_table(d, t, 400.0, ctx=ctx)
    return tbl._tbl if hasattr(tbl, "_tbl") else d.element.body.find(qn("w:tbl"))


def _twips(tbl, path):
    el = tbl.find(path)
    return None if el is None else int(el.get(qn("w:w")))


class TableEdgeAgreement(unittest.TestCase):
    def test_lead_pad_is_shared_by_the_first_column(self):
        self.assertEqual(docxout._lead_pad(_table([6.0, 6.0]).rows, 2), 6.0)
        self.assertEqual(docxout._lead_pad(_table([6.0, 9.0]).rows, 2), 6.0)
        # a differing first cell holding blocks cannot be moved by an indent
        self.assertIsNone(docxout._lead_pad(_table([6.0, 9.0], blocks_row=1).rows, 2))

    def test_first_cell_pad_becomes_the_table_default(self):
        tbl = _write(_table([6.0, 6.0], left_indent=20.0))
        ppr = tbl.find(qn("w:tblPr"))
        self.assertEqual(_twips(ppr, qn("w:tblCellMar") + "/" + qn("w:left")), 120)
        # edge stays at 20pt: Word (indent - first cell margin) and
        # LibreOffice (indent - table default) both subtract 6pt
        self.assertEqual(_twips(ppr, qn("w:tblInd")), (20 + 6) * 20)
        first = tbl.find(qn("w:tr")).find(qn("w:tc"))
        mar = first.find(qn("w:tcPr")).find(qn("w:tcMar")).find(qn("w:left"))
        self.assertEqual(int(mar.get(qn("w:w"))), 120)

    def test_a_hanging_border_is_kept(self):
        tbl = _write(_table([7.5, 7.5], left_indent=0.0, hang=7.5))
        ppr = tbl.find(qn("w:tblPr"))
        # edge 7.5pt left of the column, text on it: indent 0 - 7.5 + 7.5
        self.assertIsNone(ppr.find(qn("w:tblInd")))
        self.assertEqual(_twips(ppr, qn("w:tblCellMar") + "/" + qn("w:left")), 150)

    def test_differing_first_cells_move_the_rest_into_indent(self):
        tbl = _write(_table([6.0, 10.0], left_indent=0.0))
        rows = tbl.findall(qn("w:tr"))
        second = rows[1].find(qn("w:tc"))
        mar = second.find(qn("w:tcPr")).find(qn("w:tcMar")).find(qn("w:left"))
        self.assertEqual(int(mar.get(qn("w:w"))), 120)          # the shared 6pt
        ind = second.find(qn("w:p")).find(qn("w:pPr")).find(qn("w:ind"))
        self.assertEqual(int(ind.get(qn("w:left"))), 80)        # its other 4pt

    def test_gdocs_profile_hangs_by_its_indent(self):
        # gdocs keeps its zero default cell margin (Docs ignores it) and, since
        # WP35c, stands the table its hang left of the indent: Docs set every
        # hanging table's text the hang right of the source otherwise (c3
        # +7.00pt), and honours a negative indent (docxout._gdocs_table_hang)
        tbl = _write(_table([6.0, 6.0], left_indent=20.0, hang=3.0), profile="gdocs")
        ppr = tbl.find(qn("w:tblPr"))
        self.assertEqual(_twips(ppr, qn("w:tblCellMar") + "/" + qn("w:left")), 0)
        self.assertEqual(_twips(ppr, qn("w:tblInd")), (20 - 3) * 20)


class HangingBorderInference(unittest.TestCase):
    def test_a_table_drawn_left_of_its_text_column_records_the_hang(self):
        from reportlab.pdfgen import canvas
        from exactdoc.convert import convert
        from exactdoc import options as O
        tmp = tempfile.mkdtemp()
        pdf = os.path.join(tmp, "hang.pdf")
        c = canvas.Canvas(pdf, pagesize=(612, 792))
        c.setFont("Helvetica", 10)
        for i in range(6):                       # the text column at x=72
            c.drawString(72, 720 - 14 * i, "Body text line %d sets the column edge here." % i)
        # a 3x3 ruled grid from x=66: its text at 72, the border 6pt left
        xs, ys = (66, 200, 334, 468), (600, 582, 564, 546)
        for x in xs:
            c.line(x, ys[0], x, ys[-1])
        for y in ys:
            c.line(xs[0], y, xs[-1], y)
        for r in range(3):
            for k in range(3):
                c.drawString(xs[k] + 6, ys[r] - 13, "cell %d.%d" % (r, k))
        c.showPage()
        c.save()
        out = os.path.join(tmp, "hang.docx")
        convert(pdf, out, options=O.RAW)
        with zipfile.ZipFile(out) as z:
            xml = z.read("word/document.xml").decode("utf-8")
        import re
        tblpr = re.search(r"<w:tblPr>.*?</w:tblPr>", xml, re.S).group(0)
        mar = int(re.search(r'<w:tblCellMar>.*?<w:left w:w="(-?\d+)"', tblpr, re.S).group(1))
        ind = re.search(r'<w:tblInd w:w="(-?\d+)"', tblpr)
        ind = int(ind.group(1)) if ind else 0
        # Word and LibreOffice draw the edge at margin + ind - mar: 6pt left
        self.assertAlmostEqual((ind - mar) / 20.0, -6.0, delta=1.0)


class ParityBlanks(unittest.TestCase):
    def test_a_numbered_section_keeps_counting_where_a_restart_would_collide(self):
        secs = [HFSection(1, 1, "decimal"), HFSection(9, 1, "decimal"),
                HFSection(44, 1, "decimal")]
        out = docxout._avoid_parity_blanks(secs, 60)
        # 1..8 then 1 (8 even, 1 odd: fine); page 43 is 35, odd like 1: dropped
        self.assertEqual([s.num_start for s in out], [1, 1, None])
        self.assertEqual([s.num_start for s in secs], [1, 1, 1], "input untouched")

    def test_a_blank_lead_section_is_rebased_instead(self):
        secs = [HFSection(1, None, None, blank=True), HFSection(2, 1, "decimal")]
        out = docxout._avoid_parity_blanks(secs, 300)
        # the cover becomes page 0, so the body's 1 follows an even page
        self.assertEqual((out[0].num_start, out[0].num_fmt), (0, "decimal"))
        self.assertEqual(out[1].num_start, 1)

    def test_a_one_page_cover_is_rebased_rather_than_the_body_renumbered(self):
        # y25: a titlePg cover, then the body restarting at 1. Dropping the
        # restart would flip the odd/even header of every body page.
        secs = [HFSection(1, None, None), HFSection(2, 1, "decimal")]
        out = docxout._avoid_parity_blanks(secs, 300)
        self.assertEqual((out[0].num_start, out[1].num_start), (0, 1))

    def test_a_roman_lead_is_never_rebased_to_zero(self):
        secs = [HFSection(1, 1, "lowerRoman"), HFSection(2, 1, "decimal")]
        out = docxout._avoid_parity_blanks(secs, 30)
        self.assertEqual((out[0].num_start, out[1].num_start), (1, None))

    def test_no_collision_changes_nothing(self):
        secs = [HFSection(1, 1, "lowerRoman"), HFSection(5, 1, "decimal")]
        out = docxout._avoid_parity_blanks(secs, 40)
        self.assertEqual([(s.num_start, s.num_fmt) for s in out],
                         [(1, "lowerRoman"), (1, "decimal")])


class LocalisedFontNames(unittest.TestCase):
    SJIS_GOTHIC = "ＭＳ ゴシック".encode("cp932")

    def test_legacy_encoded_names_reach_the_family_table(self):
        name = fonts.decode_font_name(b"IKCYVK+" + self.SJIS_GOTHIC)
        self.assertEqual(name, "IKCYVK+ＭＳ ゴシック")
        fam = fonts.lookup_family(name)
        self.assertEqual((fam.cls, fam.standard, fam.ea[0]), ("sans", "Arial", "MS Gothic"))
        self.assertEqual(fonts.lookup_family(
            fonts.decode_font_name("ＭＳ Ｐゴシック".encode("cp932"))).ea[0], "MS PGothic")
        self.assertEqual(fonts.lookup_family("DFKaiShu-SB-Estd-BF").ea[0], "DFKai-SB")
        self.assertEqual(fonts.lookup_family(
            fonts.decode_font_name("新細明體".encode("big5"))).ea[0], "PMingLiU")

    def test_unknown_bytes_decode_as_before(self):
        raw = b"\xff\xfeXYZ"
        self.assertEqual(fonts.decode_font_name(raw), raw.decode("utf-8", "replace"))
        self.assertEqual(fonts.decode_font_name(b"Helvetica-Bold"), "Helvetica-Bold")
        # PyMuPDF's one-byte-per-character form
        self.assertEqual(fonts.decode_font_str(self.SJIS_GOTHIC.decode("latin-1")),
                         "ＭＳ ゴシック")
        self.assertEqual(fonts.decode_font_str("Arial"), "Arial")

    def test_every_local_name_points_at_a_table_row(self):
        for key in fonts._LOCAL_NAMES.values():
            self.assertIn(key, fonts._FAMILY_TABLE)

    def test_pdfium_reports_the_decoded_name(self):
        from exactdoc.parse_pdfium import parse_pdf
        name = "".join("#%02X" % b for b in self.SJIS_GOTHIC.replace(b" ", b""))
        content = b"BT /F1 12 Tf 72 700 Td (Hello) Tj ET"
        objs = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /TrueType /BaseFont /" + name.encode() +
            b" /FirstChar 32 /LastChar 126 /Encoding /WinAnsiEncoding >>",
            b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        ]
        buf = io.BytesIO()
        buf.write(b"%PDF-1.4\n")
        offs = []
        for i, o in enumerate(objs, 1):
            offs.append(buf.tell())
            buf.write(b"%d 0 obj\n" % i + o + b"\nendobj\n")
        xref = buf.tell()
        buf.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
        for o in offs:
            buf.write(b"%010d 00000 n \n" % o)
        buf.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
                  % (len(objs) + 1, xref))
        path = os.path.join(tempfile.mkdtemp(), "sjis.pdf")
        with open(path, "wb") as f:
            f.write(buf.getvalue())
        ir = parse_pdf(path)
        spans = [s for b in ir.pages[0].blocks for l in b.lines for s in l.spans]
        self.assertTrue(spans)
        self.assertEqual(fonts.lookup_family(spans[0].font).ea[0], "MS Gothic")


class StockContentFonts(unittest.TestCase):
    def test_non_stock_native_families_are_written_as_core_faces(self):
        for src, std in (("NotoSerif-Regular", "Cambria"),
                         ("Roboto-Medium", "Calibri"), ("Figtree-Bold", "Calibri"),
                         ("RobotoMono-Regular", "Courier New"),
                         ("SourceCodePro-Regular", "Courier New"),
                         ("Vollkorn-Regular", "Cambria")):
            self.assertEqual(fonts.writer_family(src), std, src)
            # the gdocs profile keeps the family Google Docs renders natively
            self.assertEqual(fonts.writer_family(src, profile="gdocs"),
                             fonts.map_font(src, profile="gdocs"), src)
            # line decisions keep reading map_font
            self.assertNotEqual(fonts.map_font(src), std, src)

    def test_stock_native_families_keep_their_names(self):
        for src in ("Georgia", "Verdana", "Consolas", "Tahoma", "Calibri"):
            self.assertEqual(fonts.writer_family(src), fonts.map_font(src), src)

    def test_the_run_names_the_core_face(self):
        d = docx.Document()
        r = d.add_paragraph().add_run("Hypertext Transfer Protocol")
        docxout._style_run(r, Run(text="x", font="NotoSerif-Regular", size=10,
                                  color="#000000", serif=True), "standard")
        rf = r._r.rPr.find(qn("w:rFonts"))
        self.assertEqual(rf.get(qn("w:ascii")), "Cambria")


class StockTemplateFonts(unittest.TestCase):
    def _convert(self, profile):
        from reportlab.pdfgen import canvas
        from exactdoc.convert import convert
        from exactdoc import options as O
        tmp = tempfile.mkdtemp()
        pdf = os.path.join(tmp, "t.pdf")
        c = canvas.Canvas(pdf)
        c.setFont("Helvetica", 11)
        c.drawString(72, 720, "One plain line of text.")
        c.save()
        out = os.path.join(tmp, "t.docx")
        convert(pdf, out, options=O.RAW if profile == "standard" else O.PDFIUM_GDOCS_CANDIDATE)
        with zipfile.ZipFile(out) as z:
            return {n: z.read(n).decode("utf-8") for n in
                    ("word/fontTable.xml", "word/styles.xml", "word/stylesWithEffects.xml")}

    def test_standard_profile_declares_no_courier_or_mincho(self):
        parts = self._convert("standard")
        for n, xml in parts.items():
            self.assertNotIn('"Courier"', xml, n)
            self.assertNotIn("ＭＳ 明朝", xml, n)
        self.assertIn('"Courier New"', parts["word/styles.xml"])

    def test_gdocs_profile_keeps_its_template(self):
        parts = self._convert("gdocs")
        self.assertIn('"Courier"', parts["word/fontTable.xml"])


if __name__ == "__main__":
    unittest.main()
