"""The writer's speed-ups change speed only (exactdoc/_docx_speed.py, WP20b).

Each shortcut is checked against the implementation it replaces: python-docx's
own `first_child_found_in` and `StoryPart.next_id`, and `copy.deepcopy` for
the layout copy. The corpus-wide proof -- word/*.xml byte-identical for all 95
documents in both profiles -- is in docs/evidence/refine-speed-2026-10-05.json.

    python -m unittest tests.test_docx_speed
"""
import copy
import io
import os
import random
import sys
import tempfile
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from docx import Document                                       # noqa: E402
from docx.oxml.parser import OxmlElement                         # noqa: E402
from lxml import etree                                           # noqa: E402

import exactdoc._docx_speed as S                                 # noqa: E402
from exactdoc.docxout import _copy_layout, write_docx            # noqa: E402
from exactdoc.layout import (Chunk, DocLayout, ImageEl, PageLayout, Para,  # noqa: E402
                             Run)

TAGS = ["w:rStyle", "w:rFonts", "w:b", "w:i", "w:color", "w:spacing", "w:w",
        "w:sz", "w:szCs", "w:u", "w:vertAlign", "w:rtl", "w:lang", "w:p",
        "w:tbl", "w:sectPr"]


def _png(w=40, h=30, colour=(200, 30, 30)):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (w, h), colour).save(buf, format="PNG")
    return buf.getvalue()


def _xml_parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()
                if n.startswith("word/") and n.endswith(".xml")}


class FirstChildFoundIn(unittest.TestCase):
    def test_same_element_as_python_docx_on_random_trees(self):
        rng = random.Random(20261005)
        for _ in range(3000):
            el = OxmlElement("w:rPr")
            n = rng.choice([0, 1, 2, 3, 5, 8, 12, 17, 30])
            for _k in range(n):
                if rng.random() < 0.05:
                    el.append(etree.Comment("c"))
                else:
                    el.append(OxmlElement(rng.choice(TAGS)))
            names = tuple(rng.sample(TAGS, rng.randint(1, len(TAGS))))
            self.assertIs(S._first_child_found_in(el, *names),
                          S._ORIGINAL_FIRST_CHILD(el, *names), (n, names))

    def test_installed(self):
        from docx.oxml.text.run import CT_R
        from docx.oxml.xmlchemy import BaseOxmlElement
        self.assertIs(BaseOxmlElement.first_child_found_in, S._first_child_found_in)
        self.assertIs(CT_R.clear_content, S._clear_run_content)


class ClearRunContent(unittest.TestCase):
    def test_same_children_survive_as_with_the_xpath(self):
        rng = random.Random(7)
        kids = ["w:rPr", "w:t", "w:tab", "w:br", "w:drawing", "w:t"]
        for _ in range(500):
            a = OxmlElement("w:r")
            for _k in range(rng.randint(0, 6)):
                if rng.random() < 0.15:
                    a.append(etree.Comment("note"))
                else:
                    a.append(OxmlElement(rng.choice(kids)))
            b = copy.deepcopy(a)
            S._clear_run_content(a)
            S._ORIGINAL_CLEAR_RUN(b)
            self.assertEqual(etree.tostring(a), etree.tostring(b))


class NextId(unittest.TestCase):
    def _ids(self, doc):
        return [int(x) for x in doc.part.element.xpath("//@id") if x.isdigit()]

    def test_marked_document_hands_out_the_ids_python_docx_would(self):
        data = _png()
        plain, marked = Document(), S.enable_monotonic_ids(Document())
        for doc in (plain, marked):
            for _ in range(7):
                doc.add_paragraph().add_run().add_picture(io.BytesIO(data))
        S.uninstall()
        try:
            reference = Document()
            for _ in range(7):
                reference.add_paragraph().add_run().add_picture(io.BytesIO(data))
        finally:
            S.install()
        self.assertEqual(self._ids(marked), self._ids(reference))
        self.assertEqual(self._ids(plain), self._ids(reference))

    def test_an_unmarked_document_keeps_the_full_scan(self):
        # A caller that removes a picture gets python-docx's answer (the freed
        # id again), because only exactdoc's own documents are marked.
        doc = Document()
        p = doc.add_paragraph()
        p.add_run().add_picture(io.BytesIO(_png()))
        p.add_run().add_picture(io.BytesIO(_png()))
        last = p.runs[-1]._r
        last.getparent().remove(last)
        self.assertEqual(doc.part.next_id, max(self._ids(doc)) + 1)


class BoxIndex(unittest.TestCase):
    """parse_pdfium._BoxIndex answers the white-glyph test the linear way did."""

    def test_same_answer_as_testing_every_box(self):
        from exactdoc.parse_pdfium import _BoxIndex
        rng = random.Random(1921)
        inf, nan = float("inf"), float("nan")
        for _ in range(300):
            boxes = []
            for _k in range(rng.randint(0, 60)):
                x0, y0 = rng.uniform(-50, 650), rng.uniform(-50, 850)
                bb = (x0, y0, x0 + rng.uniform(-5, 40), y0 + rng.uniform(-5, 30))
                r = rng.random()
                if r < 0.02:
                    bb = (-inf, y0, inf, y0 + 10)
                elif r < 0.04:
                    bb = (nan, y0, x0, y0 + 10)
                elif r < 0.06:
                    bb = (-1e9, -1e9, 1e9, 1e9)
                boxes.append(bb)
            idx = _BoxIndex(boxes)
            for _p in range(80):
                x, y = rng.uniform(-60, 700), rng.uniform(-60, 900)
                if rng.random() < 0.03:
                    x = rng.choice([inf, -inf, nan])
                want = any(p[0] <= x <= p[2] and p[1] <= y <= p[3] for p in boxes)
                self.assertEqual(idx.holds(x, y), want, (x, y))


class WriterOutput(unittest.TestCase):
    def _layout(self):
        els = []
        for i in range(12):
            els.append(Para(runs=[Run(text="Figure %d below." % i, font="Helvetica",
                                      size=10, color="#000000")]))
            els.append(ImageEl(data=_png(40 + i, 30, (i * 20 % 255, 80, 120)),
                               ext="png", width=120, height=90))
        return DocLayout(pages=[PageLayout(1, [Chunk(elements=els[:12])]),
                                PageLayout(2, [Chunk(elements=els[12:])])])

    def test_write_docx_is_byte_identical_with_and_without_the_shortcuts(self):
        lay = self._layout()
        with tempfile.TemporaryDirectory() as d:
            fast, slow = os.path.join(d, "fast.docx"), os.path.join(d, "slow.docx")
            write_docx(lay, fast)
            S.uninstall()
            try:
                write_docx(lay, slow)
            finally:
                S.install()
            self.assertEqual(_xml_parts(fast), _xml_parts(slow))
            ids = [int(x) for x in etree.fromstring(_xml_parts(fast)["word/document.xml"])
                   .xpath("//@id") if x.isdigit()]
            self.assertEqual(sorted(i for i in ids if i), list(range(1, 13)))

    def test_layout_copy_matches_deepcopy_and_leaves_the_source_alone(self):
        lay = self._layout()
        a, b = copy.deepcopy(lay), _copy_layout(lay)
        self.assertIsNot(b, lay)
        self.assertIsNot(b.pages[0].chunks[0].elements[0], lay.pages[0].chunks[0].elements[0])
        self.assertEqual(repr(a), repr(b))

    def test_layout_copy_falls_back_to_deepcopy(self):
        lay = self._layout()
        lay.ladder_report = {"probe": lambda: None}         # pickle cannot carry it
        b = _copy_layout(lay)
        self.assertEqual(repr(b.pages), repr(lay.pages))
        self.assertIs(b.ladder_report["probe"], lay.ladder_report["probe"])


class RunPropertiesShortcut(unittest.TestCase):
    """docxout._style_run's copied rPr is the rPr the long way builds (WP20c)."""

    FONTS = ["Helvetica", "Times-Bold", "Courier", "ABCDEF+Roboto-Regular",
             "MSGothic", "WenQuanYiZenHei", "David", "ArialMT", "NotoSansMono"]
    TEXTS = ["Plain words.", " leading and trailing ", "日本語のテキスト",
             "בפסק דין.", "שלום world", "مرحبا", "नमस्ते", "1234.", "x"]
    COLORS = ["#000000", "#1F3864", "#ff0000", None, "", "#12345", "nothex"]

    def _runs(self, rng, n):
        out = []
        for _ in range(n):
            out.append(Run(
                text=rng.choice(self.TEXTS), font=rng.choice(self.FONTS),
                size=rng.choice([6.5, 8, 9.25, 10, 10.04, 11.5, 14, 24]),
                color=rng.choice(self.COLORS), bold=rng.random() < 0.3,
                italic=rng.random() < 0.2, mono=rng.random() < 0.1,
                serif=rng.random() < 0.3, underline=rng.random() < 0.1,
                superscript=rng.random() < 0.1,
                char_spacing=rng.choice([0.0, 0.0, 0.003, -0.2, 0.35]),
                tracking=rng.choice([0.0, 0.0, 0.5]),
                width_scale=rng.choice([0.0, 0.0, 1.0, 0.97, 1.002, 1.08])))
        return out

    def _layout(self, seed):
        rng = random.Random(seed)
        els = []
        for i in range(60):
            els.append(Para(runs=self._runs(rng, rng.randint(1, 6)),
                            rtl=rng.random() < 0.2,
                            align=rng.choice(["left", "justify", "center"])))
        return DocLayout(pages=[PageLayout(1, [Chunk(elements=els)])])

    def _write(self, lay, path, profile, shortcut):
        from exactdoc import docxout
        old = docxout._RPR_SHORTCUT
        docxout._RPR_SHORTCUT = shortcut
        docxout._RPR_TEMPLATES.clear()
        try:
            write_docx(lay, path, output_profile=profile)
        finally:
            docxout._RPR_SHORTCUT = old

    def test_byte_identical_with_and_without_the_shortcut(self):
        for seed in (1, 2, 3):
            lay = self._layout(seed)
            for profile in ("standard", "gdocs"):
                with tempfile.TemporaryDirectory() as d:
                    fast, slow = os.path.join(d, "f.docx"), os.path.join(d, "s.docx")
                    self._write(lay, fast, profile, True)
                    self._write(lay, slow, profile, False)
                    self.assertEqual(_xml_parts(fast), _xml_parts(slow), (seed, profile))

    def test_templates_are_reused_and_never_the_live_element(self):
        from exactdoc import docxout
        body = Run(text="Body text.", font="Times-Roman", size=10, color="#000000")
        bold = Run(text="Heading", font="Times-Bold", size=12, color="#000000", bold=True)
        els = [Para(runs=[copy.deepcopy(bold), copy.deepcopy(body)]) for _ in range(40)]
        lay = DocLayout(pages=[PageLayout(1, [Chunk(elements=els)])])
        with tempfile.TemporaryDirectory() as d:
            docxout._RPR_TEMPLATES.clear()
            write_docx(lay, os.path.join(d, "o.docx"))
        self.assertEqual(len(docxout._RPR_TEMPLATES), 2)
        for tpl in docxout._RPR_TEMPLATES.values():
            self.assertIsNone(tpl.getparent())


class TextRunShortcut(unittest.TestCase):
    """docxout._add_text_run appends what python-docx's add_run(text) does."""

    def test_same_xml_as_add_run(self):
        from exactdoc.docxout import _add_text_run
        rng = random.Random(55)
        alphabet = ["a", "Z", " ", " ", "\t", "\n", "\r", "é", "中", "&", "<", "'"]
        doc = Document()
        for _ in range(2000):
            text = "".join(rng.choice(alphabet) for _k in range(rng.randint(0, 9)))
            p1, p2 = doc.add_paragraph(), doc.add_paragraph()
            p1.add_run("pre")
            p2.add_run("pre")
            r1, r2 = p1.add_run(text), _add_text_run(p2, text)
            self.assertIs(r2._parent, p2)
            self.assertEqual(etree.tostring(p1._p), etree.tostring(p2._p), repr(text))
            self.assertEqual(r1.text, r2.text)


class SectionBreakParagraphs(unittest.TestCase):
    """The XPath new_section uses selects what its old per-paragraph test did."""

    def test_same_paragraphs_as_the_python_scan(self):
        from docx.oxml.ns import qn
        rng = random.Random(96)
        for _ in range(200):
            doc = Document()
            body = doc.element.body
            for _k in range(rng.randint(0, 25)):
                p = doc.add_paragraph("t")._p
                shape = rng.random()
                if shape < 0.3:
                    p.get_or_add_pPr().append(OxmlElement("w:sectPr"))
                elif shape < 0.4:
                    p.append(OxmlElement("w:pPr"))           # a second pPr
                    p[-1].append(OxmlElement("w:sectPr"))
                elif shape < 0.5:
                    p.get_or_add_pPr().append(OxmlElement("w:spacing"))
            want = [p for p in body.findall(qn("w:p"))
                    if p.find(qn("w:pPr")) is not None and
                    p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None]
            self.assertEqual(body.xpath("./w:p[w:pPr[1]/w:sectPr]"), want)


if __name__ == "__main__":
    unittest.main()
