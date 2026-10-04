"""Designed pages stay editable: panels, cards, sidebars and gutter columns.

Text inside a designed region used to be rasterised or mis-flowed (design WP13):

  * a rounded panel is drawn with curves, so the parser called it 'complex'
    artwork and inference rasterised the panel with every paragraph inside it
    -- y58_ssa_statement's first page was one 1822x1574px picture, c1's three
    KPI cards one 495x59pt picture;
  * two panels side by side share baselines, and the parser joined the two
    halves of a baseline into one line that belonged to neither panel;
  * panels in two columns, a sidebar narrower than 35% of the page, a photo
    beside a name were all linearised: stacked, or interleaved line by line;
  * a CV's date column left of the main column made the main column the page's
    left margin, so every date was inlined into its role and wrapped (y44).

Each test builds the smallest PDF that shows one of these with reportlab.

    python -m unittest tests.test_designed_regions
"""
import io
import os
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from exactdoc import infer as I  # noqa: E402
from exactdoc.layout import (Cell, Chunk, ColBreak, DocLayout, FigureEl,  # noqa: E402
                             ImageEl, PageLayout, Para, RuleEl, Run, TableEl,
                             iter_paras)
from exactdoc.model import (Line, Span, TextBlock, rounded_rect_bbox)  # noqa: E402

try:
    from reportlab.pdfgen import canvas as _canvas
    from reportlab.lib.utils import ImageReader
except ImportError:                                    # pragma: no cover
    _canvas = None

W, H = 612.0, 792.0


def _y(top):
    return H - top


def _infer(path, anchored=True):
    from exactdoc.dialect import normalize
    from exactdoc.parse_pdfium import parse_pdf
    return I.infer(normalize(parse_pdf(path, keep_image_data=True)),
                   anchored=anchored)


def _elements(lay, page=1):
    return [el for ch in lay.pages[page - 1].chunks for el in ch.elements]


def _all_text(lay):
    return " ".join(p.text for p in iter_paras(lay))


def _png(w=60, h=60):
    from PIL import Image
    im = Image.new("RGB", (w, h), (30, 60, 120))
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return ImageReader(buf)


PROSE = ["You have earned enough credits to qualify for",
         "benefits. To qualify for benefits you earn credits",
         "through your work, up to four each year, and the",
         "credits you earn stay on your record."]


def _panel(c, x, top, w, h, heading, lines, fill=(0.88, 0.89, 0.91)):
    c.setFillColorRGB(*fill)
    c.roundRect(x, _y(top + h), w, h, 10, stroke=0, fill=1)
    c.setFillColorRGB(0.8, 0.1, 0.1)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(x + 9, _y(top + 16), heading)
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica", 11)
    for k, t in enumerate(lines):
        c.drawString(x + 9, _y(top + 30 + 12 * k), t)


def _statement(path):
    """y58's shape: a seal beside the masthead, rounded panels in two
    columns sharing baselines, a full-width rounded band at the foot."""
    c = _canvas.Canvas(path, pagesize=(W, H))
    c.drawImage(_png(), 36, _y(95), 58, 58)
    c.setFont("Helvetica-Bold", 30)
    c.drawString(110, _y(80), "Your Benefit Statement")
    _panel(c, 36, 150, 265, 90, "Retirement Benefits", PROSE)
    _panel(c, 36, 246, 265, 90, "Disability Benefits", PROSE)
    _panel(c, 310, 150, 266, 186, "Medicare", PROSE + PROSE)
    c.setFillColorRGB(0.88, 0.89, 0.91)
    c.roundRect(36, _y(380), 540, 28, 8, stroke=0, fill=1)
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(45, _y(368), "We base benefit estimates on current law.")
    c.showPage()
    c.save()
    return path


def _cards(path):
    c = _canvas.Canvas(path, pagesize=(W, H))
    c.setFont("Helvetica", 10)
    c.drawString(62, _y(100), "Retrieval quality degrades as the corpus grows.")
    for k, (big, small) in enumerate((("41%", "PRECISION DROP"), ("2.3x", "TAIL LATENCY"),
                                      ("$0.72", "COST PER QUERY"))):
        x = 62 + k * 166.5
        c.setFillColorRGB(0.98, 0.99, 1.0)
        c.setStrokeColorRGB(0.79, 0.83, 0.87)
        c.roundRect(x, _y(178), 157, 53, 6, stroke=1, fill=1)
        c.setFillColorRGB(0.1, 0.2, 0.4)
        c.setFont("Helvetica-Bold", 18)
        c.drawCentredString(x + 78.5, _y(150), big)
        c.setFont("Helvetica", 7)
        c.drawCentredString(x + 78.5, _y(166), small)
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica", 10)
    c.drawString(62, _y(205), "Every card above sits on one band.")
    c.showPage()
    c.save()
    return path


def _sidebar(path, shaded=False):
    """WP11's sidebar fixture: a 140pt sidebar beside a 350pt main column."""
    c = _canvas.Canvas(path, pagesize=(595, 842))
    yy = lambda t: 842 - t  # noqa: E731
    c.setFont("Helvetica-Bold", 20)
    c.drawString(40, yy(60), "Jordan Lee")
    if shaded:
        c.setFillColorRGB(0.92, 0.94, 0.97)
        c.rect(30, yy(820), 160, 740, stroke=0, fill=1)
        c.setFillColorRGB(0, 0, 0)
    top = 100
    for sec, items in (("CONTACT", ["jordan@example.com", "+1 555 0100", "example.com/jl"]),
                       ("SKILLS", ["Python, Go", "PostgreSQL", "Kubernetes", "Terraform"]),
                       ("LANGUAGES", ["English", "Spanish"])):
        c.setFont("Helvetica-Bold", 9)
        c.drawString(40, yy(top), sec)
        top += 16
        c.setFont("Helvetica", 8.5)
        for it in items:
            c.drawString(40, yy(top), it)
            top += 12
        top += 14
    top = 100
    c.setFont("Helvetica-Bold", 9)
    c.drawString(205, yy(top), "EXPERIENCE")
    top += 18
    for role, date, bullets in (
            ("Staff Engineer - Acme Corp", "2021 - 2024",
             ["Led the migration of the billing platform to an event-sourced",
              "design, cutting reconciliation time by 80%.",
              "Mentored six engineers across two teams."]),
            ("Senior Engineer - Globex", "2017 - 2021",
             ["Built the ingestion pipeline handling 2B events per day.",
              "Owned the on-call rotation and the incident review process."])):
        c.setFont("Helvetica-Bold", 9.5)
        c.drawString(205, yy(top), role)
        c.setFont("Helvetica-Oblique", 8.5)
        c.drawRightString(555, yy(top), date)
        top += 14
        c.setFont("Helvetica", 9)
        for b in bullets:
            c.drawString(205, yy(top), b)
            top += 12
        top += 10
    c.showPage()
    c.save()
    return path


def _gutter_cv(path):
    """y44's shape: dates right-aligned in a gutter beside each entry."""
    c = _canvas.Canvas(path, pagesize=(W, H))
    c.setFont("Helvetica", 24)
    c.drawString(50, _y(70), "John Doe")
    c.setFont("Helvetica", 9)
    c.drawString(50, _y(95), "San Francisco, CA   john.doe@email.com   rendercv.com")
    top = 130
    for date, role, bullets in (
            ("Sept 2018 - May 2023", "Princeton University, PhD in Computer Science",
             ["Thesis: Efficient Neural Architecture Search", "Advisor: Prof. Sanjeev Arora"]),
            ("Sept 2014 - June 2018", "Bogazici University, BS in Computer Engineering",
             ["GPA: 3.97/4.00, Valedictorian"]),
            ("June 2023 - present", "Co-Founder and CTO, Nexus AI, San Francisco, CA",
             ["Built foundation model infrastructure serving 2M+ monthly requests",
              "Raised $18M Series A led by Sequoia Capital"])):
        c.setFont("Helvetica", 10)
        c.drawRightString(167, _y(top), date)
        # rendercv sets the institution bold, and that style change is where
        # the parser's span -- and so the main column -- begins
        c.setFont("Helvetica-Bold", 10)
        c.drawString(176.5, _y(top), role)
        c.setFont("Helvetica", 10)
        top += 15
        for b in bullets:
            c.drawString(176.5, _y(top), b)
            top += 15
        top += 10
    for k in range(6):
        c.drawString(176.5, _y(top + 15 * k),
                     "Body text in the main column, line %d of the closing summary." % k)
    c.showPage()
    c.save()
    return path


def _caption(path, body=True):
    """y63's shape: a court caption, the parties | the case title, ruled
    off down the middle; the order's text runs on under it when `body`."""
    c = _canvas.Canvas(path, pagesize=(W, H))
    c.setFont("Times-Roman", 14)
    c.drawCentredString(306, _y(110), "UNITED STATES DISTRICT COURT")
    top = 160
    for text, right in (("LINDA LEE E.,", False), ("Plaintiff,", True), ("v.", False),
                        ("CAROLYN W. COLVIN,", False),
                        ("Acting Commissioner of Social Security", False),
                        ("Defendant.", True)):
        if right:
            c.drawRightString(315, _y(top), text)
        else:
            c.drawString(76, _y(top), text)
        top += 24
    c.setLineWidth(0.75)
    c.line(319, _y(146), 319, _y(top))
    c.line(72, _y(top), 319, _y(top))
    c.setFont("Times-Bold", 14)
    for k, text in enumerate(("Case No.: 23cv770-LR", "ORDER GRANTING JOINT",
                              "MOTION FOR THE AWARD AND", "PAYMENT OF ATTORNEY FEES",
                              "PURSUANT TO THE EQUAL", "[ECF No. 16]")):
        c.drawString(337, _y(160 + 24 * k), text)
    if body:
        c.setFont("Times-Roman", 14)
        for k in range(8):
            c.drawString(72, _y(top + 40 + 24 * k),
                         "The Court has reviewed the joint motion, line %d of the order." % k)
    c.showPage()
    c.save()
    return path


# ----------------------------------------------------------------- geometry
def _rrect(x0, y0, x1, y1, r):
    k = 0.5523 * r
    return [("m", (x0 + r, y0)), ("l", (x1 - r, y0)),
            ("c", (x1 - r + k, y0), (x1, y0 + r - k), (x1, y0 + r)),
            ("l", (x1, y1 - r)),
            ("c", (x1, y1 - r + k), (x1 - r + k, y1), (x1 - r, y1)),
            ("l", (x0 + r, y1)),
            ("c", (x0 + r - k, y1), (x0, y1 - r + k), (x0, y1 - r)),
            ("l", (x0, y0 + r)),
            ("c", (x0, y0 + r - k), (x0 + r - k, y0), (x0 + r, y0))]


class RoundedRectangles(unittest.TestCase):
    def test_a_panel_with_rounded_corners_is_a_box(self):
        self.assertEqual(rounded_rect_bbox(_rrect(36, 287, 301, 470, 10)),
                         (36, 287, 301, 470))

    def test_a_pill_is_a_box(self):
        self.assertIsNotNone(rounded_rect_bbox(_rrect(0, 0, 130, 14, 7)))

    def test_a_circle_is_not(self):
        # four quarter arcs and no straight edge: pi/4 of its box
        self.assertIsNone(rounded_rect_bbox(_rrect(0, 0, 40, 40, 20)))

    def test_a_diagonal_edge_is_not(self):
        segs = _rrect(0, 0, 100, 40, 5)
        segs[1] = ("l", (90, 6))          # a slanted top edge
        self.assertIsNone(rounded_rect_bbox(segs))

    def test_a_glyph_sized_one_stays_a_glyph(self):
        # dialect reads bullets and markers from curves this small
        self.assertIsNone(rounded_rect_bbox(_rrect(0, 0, 8, 8, 2)))

    def test_two_subpaths_are_not(self):
        segs = _rrect(0, 0, 100, 40, 5) + [("m", (10, 10))] + _rrect(10, 10, 90, 30, 3)[1:]
        self.assertIsNone(rounded_rect_bbox(segs))


def _span(text, x0, top, size=11.0, x1=None, bold=False, font="Helvetica"):
    x1 = x1 if x1 is not None else x0 + 0.5 * size * len(text)
    return Span(text=text, font=font, size=size, color="#000000", bold=bold,
                italic=False, mono=False, serif=False, superscript=False,
                bbox=(x0, top, x1, top + 1.2 * size), origin=(x0, top + 0.9 * size))


def _line(*spans):
    return Line(spans=list(spans), bbox=(min(s.bbox[0] for s in spans),
                                          min(s.bbox[1] for s in spans),
                                          max(s.bbox[2] for s in spans),
                                          max(s.bbox[3] for s in spans)))


class LinesAcrossPanelEdges(unittest.TestCase):
    def _draws(self):
        # the two panels' boxes, as the leftover-box pass will build them
        return [(35.9, 474.0, 301.4, 543.1), (310.4, 471.5, 575.9, 722.5)]

    def _joined(self):
        ln = _line(_span("You have earned enough credits to qualif", 45.0, 490.7, x1=292.0),
                   _span("You have enough credits to qualify for M", 319.4, 490.7, x1=562.3))
        return ln, TextBlock(lines=[ln], bbox=ln.bbox)

    def test_a_line_joined_across_two_panels_is_cut_at_the_gutter(self):
        ln, blk = self._joined()
        self.assertEqual(len(I._split_lines_at_box_edges([blk], self._draws())), 1)
        self.assertEqual([l.text for l in blk.lines],
                         ["You have earned enough credits to qualif",
                          "You have enough credits to qualify for M"])

    def test_a_line_no_box_takes_a_piece_of_stays_whole(self):
        # pieces that would both stay in the flow are two lines on one
        # baseline, which the flow stacks
        ln = _line(_span("left column text ending here", 45.0, 600.0, x1=200.0),
                   _span("right column text", 240.0, 600.0, x1=330.0))
        blk = TextBlock(lines=[ln], bbox=ln.bbox)
        self.assertEqual(I._split_lines_at_box_edges([blk], [(220.0, 100.0, 400.0, 150.0)]), [])

    def test_a_consumed_line_is_never_cut(self):
        ln, blk = self._joined()
        self.assertEqual(I._split_lines_at_box_edges([blk], self._draws(), {id(ln)}), [])

    def test_a_word_touching_a_panel_edge_stays_whole(self):
        ln = _line(_span("credits to qualify for", 200.0, 490.7, x1=299.0),
                   _span("disability", 302.0, 490.7, x1=340.0))
        blk = TextBlock(lines=[ln], bbox=ln.bbox)
        self.assertEqual(I._split_lines_at_box_edges([blk], self._draws()), [])

    def test_a_cut_no_region_claimed_is_joined_back(self):
        # the rectangles were boxes-to-be; if no box took a piece (a figure
        # absorbed the rect, a table claimed it), the line is put back whole
        ln, blk = self._joined()
        cuts = I._split_lines_at_box_edges([blk], self._draws())
        I._restore_uncut(cuts, consumed=set())
        self.assertEqual([l.text for l in blk.lines], [ln.text])
        self.assertIs(blk.lines[0], ln)

    def test_a_cut_whose_pieces_all_found_regions_stays_cut(self):
        ln, blk = self._joined()
        cuts = I._split_lines_at_box_edges([blk], self._draws())
        I._restore_uncut(cuts, consumed={id(l) for l in blk.lines})
        self.assertEqual(len(blk.lines), 2)

    def test_a_cut_with_a_piece_left_in_the_flow_is_joined_back(self):
        # body text running past a box: re-paragraphing it around the box
        # cost y60 a page, so the line stays whole
        ln, blk = self._joined()
        cuts = I._split_lines_at_box_edges([blk], self._draws())
        I._restore_uncut(cuts, consumed={id(blk.lines[0])})
        self.assertEqual([l.text for l in blk.lines], [ln.text])

class RowOfItems(unittest.TestCase):
    def _items(self, lines):
        out = []
        for ln in lines:
            out.append(("blk", ln.bbox, TextBlock(lines=[ln], bbox=ln.bbox)))
        return out

    def test_items_on_one_baseline_are_one_tabbed_row(self):
        # y44's contact strip: five separate items 14pt apart
        xs = [(50.4, 139.2, "San Francisco, CA"), (153.3, 259.6, "john.doe@email.com"),
              (273.7, 347.9, "rendercv.com"), (362.0, 413.4, "rendercv"),
              (427.6, 480.1, "rendercv")]
        lines = [_line(_span(t, x, 88.0, size=10.0, x1=x1)) for x, x1, t in xs]
        flow = I._to_flow(self._items(lines), 47.6, 561.6)
        paras = [e for e in flow if isinstance(e, Para)]
        self.assertEqual(len(paras), 1)
        self.assertEqual(paras[0].text.count("\t"), 4)
        self.assertEqual(len(paras[0].tab_stops), 4)

    def test_two_items_are_not_a_row(self):
        lines = [_line(_span("left", 50.0, 88.0, size=10.0, x1=70.0)),
                 _line(_span("right", 300.0, 88.0, size=10.0, x1=325.0))]
        flow = I._to_flow(self._items(lines), 47.6, 561.6)
        self.assertEqual(len([e for e in flow if isinstance(e, Para)]), 2)

    def test_prose_lines_are_not_items(self):
        t = "a line of running prose with many words in it"
        lines = [_line(_span(t, x, 88.0, size=10.0, x1=x + 150.0)) for x in (50.0, 210.0, 370.0)]
        flow = I._to_flow(self._items(lines), 47.6, 561.6)
        self.assertEqual(len([e for e in flow if isinstance(e, Para)]), 3)

class ForcedBreaks(unittest.TestCase):
    def test_a_heading_line_over_its_text_is_its_own_paragraph(self):
        head = _line(_span("Retirement Benefits", 45.0, 300.0, size=12.0, x1=155.0, bold=True))
        body = [_line(_span(t, 45.0, 314.0 + 12 * k, x1=290.0)) for k, t in enumerate(PROSE)]
        merged = I.paras_from_line_list([head] + body, 45.0, 297.0)
        split = I.paras_from_line_list([head] + body, 45.0, 297.0, forced=297.0)
        self.assertEqual(len(merged), 1)
        self.assertEqual([p.text for p in split][0], "Retirement Benefits")
        self.assertEqual(len(split), 2)
        self.assertTrue(split[1]._forced)

    def test_a_wrapped_line_is_not_a_break(self):
        a = _line(_span("a line that wrapped because the next word did", 45.0, 300.0, x1=280.0))
        b = _line(_span("extraordinarily not fit", 45.0, 312.0, x1=160.0))
        self.assertFalse(I._forced_break(a, b, 297.0))

    def test_a_hyphenated_end_is_a_wrap(self):
        a = _line(_span("a line ending in a hyphen-", 45.0, 300.0, x1=150.0))
        b = _line(_span("ated word", 45.0, 312.0, x1=100.0))
        self.assertFalse(I._forced_break(a, b, 297.0))

    def _caption(self):
        # y63's caption: the parties flush left, "Plaintiff," and
        # "Defendant." flush right against the caption's rule
        rows = [("LINDA LEE E.,1", 75.6, 171.9), ("Plaintiff,", 266.2, 315.6),
                ("v.", 75.6, 86.2), ("CAROLYN W. COLVIN,2", 75.6, 228.0),
                ("Acting Commissioner of Social Security", 75.6, 302.6),
                ("Defendant.", 253.9, 315.7)]
        return [_line(_span(t, x0, 303.0 + 24.0 * k, size=14.04, x1=x1))
                for k, (t, x0, x1) in enumerate(rows)]

    def test_flush_right_lines_in_a_flush_left_column_are_lines_of_their_own(self):
        lines = self._caption()
        edge = I._text_column_edge(lines)
        self.assertEqual(edge, 302.6)               # the flush-left lines' edge
        paras = I.paras_from_line_list(lines, 75.6, 315.7, forced=edge)
        self.assertEqual([p.text for p in paras],
                         ["LINDA LEE E.,1", "Plaintiff,", "v.", "CAROLYN W. COLVIN,2",
                          "Acting Commissioner of Social Security", "Defendant."])

    def test_two_columns_in_one_frame_are_still_not_one_column(self):
        # y59's shape: the lines off the left edge are a column, not lines
        # set flush right, and nothing but the heading rule may break
        left = [_line(_span("left column text runs on", 50.0, 300.0 + 12 * k, x1=240.0))
                for k in range(4)]
        right = [_line(_span("right column", 300.0, 300.0 + 12 * k, x1=420.0 + 20 * k))
                 for k in range(4)]
        self.assertIsNone(I._flush_right_edge(left + right))
        self.assertEqual(I._text_column_edge(left + right), float("-inf"))

    def test_no_line_passes_an_ordinary_columns_edge(self):
        a = _line(_span("a full line of text", 45.0, 300.0, x1=297.0))
        b = _line(_span("and its next", 45.0, 312.0, x1=200.0))
        self.assertFalse(I._forced_break(a, b, 297.0))
        self.assertFalse(I._forced_break(a, b, float("-inf")))


@unittest.skipIf(_canvas is None, "reportlab is not installed")
class DesignedPagesEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._dir = tempfile.TemporaryDirectory()
        d = cls._dir.name
        cls.statement = _infer(_statement(os.path.join(d, "statement.pdf")))
        # the Google Docs profile keeps graphics in the flow (no `anchored`)
        cls.statement_flow = _infer(os.path.join(d, "statement.pdf"),
                                    anchored=False)
        cls.cards = _infer(_cards(os.path.join(d, "cards.pdf")))
        cls.sidebar = _infer(_sidebar(os.path.join(d, "sidebar.pdf")))
        cls.shaded = _infer(_sidebar(os.path.join(d, "shaded.pdf"), shaded=True))
        cls.cv = _infer(_gutter_cv(os.path.join(d, "cv.pdf")))
        cls.caption = _infer(_caption(os.path.join(d, "caption.pdf")))
        cls.caption_alone = _infer(_caption(os.path.join(d, "caption_alone.pdf"),
                                            body=False))

    @classmethod
    def tearDownClass(cls):
        cls._dir.cleanup()

    def test_rounded_panels_are_editable_boxes_not_pictures(self):
        els = _elements(self.statement)
        self.assertFalse([e for e in els if isinstance(e, FigureEl)])
        text = _all_text(self.statement)
        for heading in ("Retirement Benefits", "Disability Benefits", "Medicare",
                        "We base benefit estimates"):
            self.assertIn(heading, text)

    def test_panels_in_two_columns_are_a_two_column_section(self):
        chunks = self.statement.pages[0].chunks
        cols = [ch for ch in chunks if ch.n_cols == 2]
        self.assertEqual(len(cols), 1)
        els = cols[0].elements
        k = next(i for i, e in enumerate(els) if isinstance(e, ColBreak))
        left = [e for e in els[:k] if isinstance(e, TableEl)]
        right = [e for e in els[k + 1:] if isinstance(e, TableEl)]
        self.assertEqual(len(left), 2)
        self.assertEqual(len(right), 1)
        # the panels keep their own heading as a paragraph of its own
        self.assertEqual(left[0].rows[0][0].paras[0].text, "Retirement Benefits")

    def test_a_seal_set_into_the_top_margin_is_anchored_where_it_is(self):
        # The seal starts 16pt above the masthead's text, in the top margin,
        # which a flow cannot reach: where the profile positions graphics,
        # WP15's `_float_backgrounds` anchors it, and the masthead flows.
        pg = self.statement.pages[0]
        self.assertEqual([type(f.el) for f in pg.floats], [ImageEl])
        first = pg.chunks[0].elements[0]
        self.assertIsInstance(first, Para)
        self.assertIn("Your Benefit Statement", first.text)

    def test_a_seal_beside_the_masthead_is_a_layout_row(self):
        # ...and where it keeps the flow (gdocs), the seal and the masthead
        # are one layout row.
        first = self.statement_flow.pages[0].chunks[0].elements[0]
        self.assertIsInstance(first, TableEl)
        self.assertEqual(first.role, "layout")
        cells = first.rows[0]
        self.assertIsInstance(cells[0].blocks[0], ImageEl)
        self.assertIn("Your Benefit Statement", cells[1].paras[0].text)

    def test_a_row_of_cards_is_one_table(self):
        tables = [e for e in _elements(self.cards) if isinstance(e, TableEl)]
        self.assertEqual(len(tables), 1, [t.role for t in tables])
        t = tables[0]
        self.assertEqual(t.role, "cards")
        self.assertEqual(len(t.col_widths), 5)      # three cards, two gutters
        texts = [" / ".join(p.text for p in c.paras) for c in t.rows[0]]
        self.assertEqual([x for x in texts if x],
                         ["41% / PRECISION DROP", "2.3x / TAIL LATENCY",
                          "$0.72 / COST PER QUERY"])

    def test_a_sidebar_narrower_than_the_two_column_bar_is_a_column(self):
        tables = [e for e in _elements(self.sidebar)
                  if isinstance(e, TableEl) and e.role == "layout"]
        self.assertEqual(len(tables), 1)
        side, main = tables[0].rows[0]
        side_text = " ".join(p.text for p in side.paras)
        main_text = " ".join(p.text for p in main.paras)
        self.assertIn("CONTACT", side_text)
        self.assertIn("Terraform", side_text)
        self.assertIn("EXPERIENCE", main_text)
        self.assertNotIn("Terraform", main_text)
        self.assertNotIn("Staff Engineer", side_text)
        # each contact line is its own paragraph, not a run-on line
        self.assertIn("jordan@example.com", [p.text for p in side.paras])

    def test_a_shaded_sidebar_is_a_column_not_a_740pt_box_above_the_page(self):
        tables = [e for e in _elements(self.shaded)
                  if isinstance(e, TableEl) and e.role == "layout"]
        self.assertEqual(len(tables), 1)
        side, main = tables[0].rows[0]
        boxes = [b for b in side.blocks if isinstance(b, TableEl)]
        self.assertEqual(len(boxes), 1)
        self.assertEqual(boxes[0].rows[0][0].shading, "#ebf0f7")
        self.assertIn("EXPERIENCE", " ".join(p.text for p in main.paras))

    def test_a_ruled_caption_over_the_pages_text_is_a_box_of_two_cells(self):
        els = _elements(self.caption)
        tables = [e for e in els if isinstance(e, TableEl) and e.role == "layout"]
        self.assertEqual(len(tables), 1)
        self.assertFalse([ch for ch in self.caption.pages[0].chunks if ch.n_cols > 1])
        parties, title = tables[0].rows[0]
        texts = [p.text for p in parties.paras]
        # the flush-right labels stay lines of their own
        self.assertIn("Plaintiff,", texts)
        self.assertIn("Defendant.", texts)
        self.assertIn("v.", texts)
        self.assertIn("Case No.: 23cv770-LR", " ".join(p.text for p in title.paras))

    def test_a_ruled_band_that_ends_the_page_stays_two_columns(self):
        self.assertTrue([ch for ch in self.caption_alone.pages[0].chunks
                         if ch.n_cols == 2])

    def test_gutter_dates_hang_beside_their_entries(self):
        self.assertLess(self.cv.margin_l, 55.0)
        dated = [p for p in iter_paras(self.cv) if p.text.startswith("\tSept 2018")]
        self.assertEqual(len(dated), 1)
        p = dated[0]
        self.assertAlmostEqual(p.left_indent, 176.5 - self.cv.margin_l, delta=1.0)
        self.assertAlmostEqual(p.first_indent, -p.left_indent, delta=0.1)
        self.assertEqual([a for _x, a in p.tab_stops], ["right", "left"])
        self.assertIn("Princeton University", p.text.split("\t")[2])
        # the main column keeps its own indent from the new margin
        bullet = next(q for q in iter_paras(self.cv) if "Thesis" in q.text)
        self.assertAlmostEqual(bullet.left_indent, 176.5 - self.cv.margin_l, delta=1.0)


class SideEvidenceRefuses(unittest.TestCase):
    def _blk(self, lines):
        bb = (min(l.bbox[0] for l in lines), min(l.bbox[1] for l in lines),
              max(l.bbox[2] for l in lines), max(l.bbox[3] for l in lines))
        return ("blk", bb, TextBlock(lines=lines, bbox=bb))

    def test_label_and_value_rows_are_a_form_not_two_columns(self):
        left = [self._blk([_line(_span("Label %d" % k, 40.0, 100.0 + 14 * k, x1=120.0))])
                for k in range(8)]
        right = [self._blk([_line(_span("value number %d of the form" % k, 200.0,
                                        100.0 + 14 * k, x1=480.0))]) for k in range(8)]
        from exactdoc.model import PageIR
        page = PageIR(number=1, width=595.0, height=842.0)
        self.assertIsNone(I._side_evidence(left, right, 120.0, 200.0, page, 515.0))

    def test_line_numbers_beside_a_signature_are_not_a_column(self):
        nums = self._blk([_line(_span(str(k), 44.0, 280.0 + 24 * k, x1=58.0))
                          for k in range(10, 28)])
        img = ImageEl(data=b"", ext="png", width=235, height=81)
        img._bbox = (300.0, 600.0, 535.0, 681.0)
        from exactdoc.model import PageIR
        page = PageIR(number=5, width=612.0, height=792.0)
        self.assertIsNone(I._side_evidence([nums], [("el", img._bbox, img)],
                                           58.0, 300.0, page, 531.0))


class LayoutTableWriter(unittest.TestCase):
    def _lay(self):
        lay = DocLayout(page_w=595.0, page_h=842.0, margin_l=40.0, margin_r=40.0,
                        margin_t=40.0, margin_b=40.0)
        mk = lambda t: Para(runs=[Run(text=t, font="Helvetica", size=9.0,  # noqa: E731
                                      color="#000000")], leading=11.0, src_lines=1)
        box = TableEl(rows=[[Cell(paras=[mk("CONTACT"), mk("jordan@example.com")],
                                  shading="#ebf0f7", pad=(6.0, 10.0, 6.0, 4.0))]],
                      col_widths=[160.0], role="box", bbox=(30.0, 80.0, 190.0, 300.0))
        side = Cell(pad=(0.0, 0.0, 0.0, 0.0))
        side.blocks = [mk("Jordan Lee"), box]
        side.paras = [side.blocks[0]]
        rule = RuleEl(width_pct=100.0, thickness=0.6, color="#000000", length=300.0)
        main = Cell(pad=(0.0, 0.0, 0.0, 0.0))
        main.blocks = [mk("EXPERIENCE"), rule, mk("Staff Engineer")]
        main.paras = [main.blocks[0], main.blocks[2]]
        t = TableEl(rows=[[side, main]], col_widths=[175.0, 350.0], row_heights=[300.0],
                    role="layout", bbox=(30.0, 40.0, 555.0, 340.0))
        t.left_indent = -10.0
        pg = PageLayout(number=1, chunks=[Chunk(n_cols=1, elements=[t])])
        lay.pages = [pg]
        return lay

    def test_a_layout_cell_holds_its_column_in_order(self):
        from docx import Document
        from docx.oxml.ns import qn
        from exactdoc.docxout import write_docx
        with tempfile.TemporaryDirectory() as d:
            out = os.path.join(d, "layout.docx")
            write_docx(self._lay(), out)
            doc = Document(out)
            self.assertEqual(len(doc.tables), 1)
            tbl = doc.tables[0]
            ind = tbl._tbl.tblPr.find(qn("w:tblInd"))
            self.assertEqual(ind.get(qn("w:w")), "-200")       # -10pt, into the margin
            side, main = tbl.rows[0].cells
            self.assertEqual(side.paragraphs[0].text, "Jordan Lee")
            self.assertEqual(len(side.tables), 1)               # the shaded box, nested
            self.assertEqual(side.tables[0].rows[0].cells[0].paragraphs[0].text, "CONTACT")
            self.assertEqual(side._tc[-1].tag, qn("w:p"))       # a cell ends in a paragraph
            self.assertEqual([p.text for p in main.paragraphs],
                             ["EXPERIENCE", "", "Staff Engineer"])
            trh = tbl.rows[0]._tr.trPr.find(qn("w:trHeight"))
            self.assertEqual(trh.get(qn("w:hRule")), "atLeast")

    def test_layout_cell_paragraphs_are_document_paragraphs(self):
        texts = [p.text for p in iter_paras(self._lay())]
        self.assertEqual(texts, ["Jordan Lee", "CONTACT", "jordan@example.com",
                                 "EXPERIENCE", "Staff Engineer"])


class LayoutRowPin(unittest.TestCase):
    """A layout row cannot split: its pin keeps two line pitches and
    LAYOUT_ROW_RESERVE_PT clear of the page foot (live Docs, 2026-10-04: a
    778pt row against a 786pt body turned one page into three)."""

    def _table(self):
        lay = LayoutTableWriter()._lay()
        return lay, _elements(lay)[0]

    def test_a_row_well_short_of_the_foot_keeps_its_region_height(self):
        lay, t = self._table()                     # body foot 802, top 40
        self.assertEqual(I._layout_row_pin(t, 40.0, lay), 300.0)

    def test_a_row_near_the_foot_is_capped_two_lines_and_reserve_clear(self):
        lay, t = self._table()
        t.row_heights = [780.0]                    # 40 + 780 runs past 802
        pin = I._layout_row_pin(t, 40.0, lay)
        # foot 802 - 2 x 11pt leading - reserve, measured from the row's top
        self.assertAlmostEqual(pin, 802.0 - 22.0 - I.LAYOUT_ROW_RESERVE_PT - 40.0)
        self.assertLessEqual(40.0 + pin + 2 * 11.0 + I.LAYOUT_ROW_RESERVE_PT,
                             lay.page_h - lay.margin_b)

    def test_a_row_with_no_room_left_is_not_pinned(self):
        lay, t = self._table()
        self.assertIsNone(I._layout_row_pin(t, 790.0, lay))

    def test_a_nested_box_drawn_to_the_foot_gives_up_its_bottom_pad(self):
        lay, t = self._table()
        box = t.rows[0][0].blocks[1]
        box.bbox = (30.0, 80.0, 190.0, 790.0)      # 14pt past the capped foot
        box.rows[0][0].pad = (6.0, 10.0, 30.0, 4.0)
        I._layout_row_pin(t, 40.0, lay)
        self.assertEqual(box.rows[0][0].pad[2], 16.0)

    def test_the_shaded_sidebar_row_stays_clear_of_the_page_foot(self):
        with tempfile.TemporaryDirectory() as d:
            lay = _infer(_sidebar(os.path.join(d, "shaded.pdf"), shaded=True))
        t = next(e for e in _elements(lay)
                 if isinstance(e, TableEl) and e.role == "layout")
        lead = max(p.leading for c in t.rows[0] for p in I._cell_all_paras(c)
                   if p.leading)
        self.assertIsNotNone(t.row_heights[0])
        self.assertLessEqual(t.bbox[1] + t.row_heights[0] + 2 * lead,
                             lay.page_h - lay.margin_b - I.LAYOUT_ROW_RESERVE_PT + 0.1)


class GdocsShadedBox(unittest.TestCase):
    """Under gdocs a box is bordered paragraphs; a filled box keeps its fill."""

    def _write(self, shading, borders):
        from docx import Document
        from docx.oxml.ns import qn
        from exactdoc.docxout import WriteCtx, _write_box_paragraphs
        mk = lambda t, top: Para(runs=[Run(text=t, font="Helvetica", size=9.0,  # noqa: E731
                                           color="#000000")], leading=11.0,
                                 bbox=(40.0, top, 180.0, top + 9.0))
        cell = Cell(paras=[mk("CONTACT", 90.0), mk("jordan@example.com", 102.0)],
                    shading=shading, borders=borders, pad=(6.0, 10.0, 6.0, 4.0))
        t = TableEl(rows=[[cell]], col_widths=[160.0], role="box",
                    bbox=(30.0, 80.0, 190.0, 120.0))
        doc = Document()
        _write_box_paragraphs(doc, t, 482.0, WriteCtx(output_profile="gdocs"))
        out = []
        for p in doc.paragraphs:
            ppr = p._p.find(qn("w:pPr"))
            if ppr is None or ppr.find(qn("w:pBdr")) is None:
                continue
            bd = ppr.find(qn("w:pBdr"))
            shd = ppr.find(qn("w:shd"))
            out.append(({c.tag.split("}")[1]: c.get(qn("w:color")) for c in bd},
                        None if shd is None else (shd.get(qn("w:val")), shd.get(qn("w:fill"))),
                        bd.getnext() is shd))
        return out

    def test_a_filled_unstroked_box_is_shaded_with_rails_in_its_fill(self):
        paras = self._write("#ebf0f7", {})
        self.assertEqual(len(paras), 2)
        for colours, shd, after_bdr in paras:
            self.assertEqual(shd, ("clear", "EBF0F7"))
            self.assertTrue(after_bdr)             # schema order: pBdr, shd
            self.assertEqual(set(colours.values()), {"EBF0F7"})
        self.assertEqual(set(paras[0][0]), {"top", "left", "right"})
        self.assertEqual(set(paras[1][0]), {"left", "bottom", "right"})

    def test_a_stroked_filled_box_keeps_its_stroke(self):
        edge = (0.75, "#1f2937")
        paras = self._write("#fef3c7", {"left": edge, "right": edge,
                                        "top": edge, "bottom": edge})
        for colours, shd, _ in paras:
            self.assertEqual(shd, ("clear", "FEF3C7"))
            self.assertEqual(set(colours.values()), {"1F2937"})

    def test_an_unfilled_box_has_no_shading_and_dark_rails(self):
        paras = self._write(None, {})
        for colours, shd, _ in paras:
            self.assertIsNone(shd)
            self.assertEqual(set(colours.values()), {"333333"})


if __name__ == "__main__":
    unittest.main()
