"""Named Heading styles for Google Docs' outline (the sidebar ask).

`w:outlineLvl` is what Word's navigation pane reads; Google Docs' outline
sidebar and style dropdown key on the paragraph STYLE. Converted documents
carried only the outline level and showed "Normal text" everywhere. The
writer now names `Heading N` on heading paragraphs and strips the stock
style definitions down to metadata, so everything visual still comes from
direct formatting and nothing moves.
"""
import io
import unittest
import zipfile

from docx import Document

from exactdoc.layout import Chunk, DocLayout, PageLayout, Para, Run


def _heading_para(text, level, size=16.0):
    p = Para(runs=[Run(text=text, font="Georgia", size=size, color="#000000",
                       bold=True, italic=False, mono=False, serif=True)],
             heading=level)
    p.leading = size * 1.2
    p._b1 = 100.0
    p._size1 = size
    p._vis_lines = 1
    p.bbox = (57.0, 90.0, 500.0, 106.0)
    return p


def _write(layout):
    from exactdoc.docxout import write_docx
    buf = io.BytesIO()
    import tempfile, os
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    try:
        write_docx(layout, path, dpi=240, output_profile="standard")
        with open(path, "rb") as fh:
            buf.write(fh.read())
    finally:
        os.unlink(path)
    return buf.getvalue()


def _layout(*paras):
    page = PageLayout(number=1, chunks=[Chunk(elements=list(paras))])
    lay = DocLayout(src_path="x.pdf")
    lay.page_w, lay.page_h = 595.0, 842.0
    lay.margin_l = lay.margin_r = 57.0
    lay.margin_t, lay.margin_b = 57.0, 57.0
    lay.pages = [page]
    return lay


class OutlineStyles(unittest.TestCase):
    def test_heading_paragraph_carries_named_style(self):
        blob = _write(_layout(_heading_para("A Title", 1)))
        z = zipfile.ZipFile(io.BytesIO(blob))
        doc = z.read("word/document.xml").decode("utf-8")
        self.assertIn('<w:pStyle w:val="Heading1"/>', doc)
        self.assertIn('<w:outlineLvl w:val="0"/>', doc)

    def test_stock_heading_style_has_no_visual_payload(self):
        blob = _write(_layout(_heading_para("A Title", 2)))
        z = zipfile.ZipFile(io.BytesIO(blob))
        styles = z.read("word/styles.xml").decode("utf-8")
        import re
        m = re.search(r'<w:style [^>]*w:styleId="Heading2".*?</w:style>',
                      styles, re.S)
        self.assertIsNotNone(m, "Heading2 style must exist")
        seg = m.group(0)
        # no font, size or colour of its own: every visual property the
        # reader sees must come from the paragraph's direct formatting
        self.assertNotIn("<w:rPr>", seg)
        # but the paragraph properties are EXPLICITLY zero, not absent --
        # LibreOffice supplies its own built-in spacing for a silent
        # "heading 1" definition, which moved a gated document
        self.assertIn('<w:spacing w:before="0" w:after="0" w:line="240"'
                      ' w:lineRule="auto"/>', seg)
        self.assertIn("w:name", seg)      # still a heading to every reader

    def test_body_paragraph_is_unstyled(self):
        body = _heading_para("plain body", 0)
        blob = _write(_layout(body))
        doc = zipfile.ZipFile(io.BytesIO(blob)) \
            .read("word/document.xml").decode("utf-8")
        self.assertNotIn('w:pStyle w:val="Heading', doc)

    def test_document_opens_and_reports_heading(self):
        blob = _write(_layout(_heading_para("A Title", 1),
                              _heading_para("Section", 2)))
        d = Document(io.BytesIO(blob))
        self.assertEqual(d.paragraphs[0].style.name, "Heading 1")
        self.assertEqual(d.paragraphs[1].style.name, "Heading 2")


if __name__ == "__main__":
    unittest.main()
