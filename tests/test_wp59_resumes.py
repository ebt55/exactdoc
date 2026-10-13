"""WP59: resume layouts that holdout-2 failed on.

Each mechanism is reproduced by a synthetic page (tests/wp59_fixtures.py)
beside the look-alike the converter must keep:

R1  Word's page colour -- one fill over the whole sheet of every page -- was
    read as a cover band and a 792pt running header strip: every page's text
    went into the header and the body section's top margin passed the page,
    which Word cannot export (`infer._page_colour`, `detect_hf`). It is now
    the document background (`w:background`), and its page border
    (`infer._strip_page_border`) is no box. Backstops: a running strip past
    mid-page is no header (`detect_hf`), and no section is written whose
    margins leave less than a body line (`docxout._config_section`).

R2  a resume's dates at a right tab beside a skills list set flush right
    beside a coursework list: the run of right-aligned lines tripped WP49's
    sidebar guard, and the dates went behind a column break
    (`infer._two_sided_run`); the fields' edge is the band of right edges
    holding most right-hand lines, not the widest line.
R3a fields at a LEFT tab -- all starting at one x, ragged at the end -- were
    never paired, and label, field and the next line welded into one
    paragraph (`infer._left_tab_pairs`).
Text font boxes: a text face whose glyph boxes are several times its size
    (Word's Cambria) put every line far above its baseline, and box-top order
    set entries above their headings (`infer._true_text_boxes`).

Every WP59 reading is the standard profile's: `infer(tab_layouts=False)` --
what the gdocs profile passes -- reproduces the earlier layout.
"""
import os
import re
import sys
import tempfile
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from exactdoc import docxout, infer as I  # noqa: E402
from exactdoc.layout import Para, TableEl, iter_paras  # noqa: E402
from exactdoc.model import DocIR, DrawCmd, PageIR  # noqa: E402

import wp59_fixtures as F  # noqa: E402


def _parse(path):
    from exactdoc.dialect import normalize
    from exactdoc.parse_pdfium import parse_pdf
    return normalize(parse_pdf(path, keep_image_data=True))


def _infer(path, **kw):
    return I.infer(_parse(path), **kw)


def _shape(lay):
    out = [("margins", round(lay.margin_t, 1), round(lay.margin_b, 1)),
           ("cover", lay.cover_band is not None), ("header", lay.header_default is not None)]
    for pl in lay.pages:
        for ch in pl.chunks:
            out.append(("chunk", ch.n_cols))
            for el in ch.elements:
                out.append(("para", el.text) if isinstance(el, Para) else (type(el).__name__,))
    return out


def _docx(lay, path):
    docxout.write_docx(lay, path, output_profile="standard")
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n).decode("utf-8") for n in z.namelist()
                if n.startswith("word/") and n.endswith(".xml")}


def _sections(xml):
    """[(page_h_tw, top_tw, bottom_tw)] for every w:sectPr."""
    out = []
    for sect in re.findall(r"<w:sectPr.*?</w:sectPr>", xml, re.S):
        h = int(re.search(r'<w:pgSz[^>]*w:h="(\d+)"', sect).group(1))
        mt = int(re.search(r'<w:pgMar[^>]*w:top="(-?\d+)"', sect).group(1))
        mb = int(re.search(r'<w:pgMar[^>]*w:bottom="(-?\d+)"', sect).group(1))
        out.append((h, mt, mb))
    return out


def _body_text(lay):
    return " ".join(p.text for p in iter_paras(lay))


# ------------------------------------------------------------- unit level
def _fill(bbox, colour="#dbe4f0", shape="rect"):
    return DrawCmd(kind="fill", shape=shape, bbox=bbox, fill=colour, stroke=None,
                   width=0.0, opacity=1.0, n_items=5)


def _ir(fills_per_page):
    return DocIR(path="synthetic.pdf", pages=[PageIR(i + 1, 612.0, 792.0, blocks=[], drawings=list(f))
                        for i, f in enumerate(fills_per_page)])


SHEET = (0.0, 0.0, 612.1, 792.1)


class PageColourUnit(unittest.TestCase):
    def test_the_same_sheet_on_every_later_page_is_the_page_colour(self):
        self.assertEqual(I._page_colour(_ir([[_fill(SHEET)]] * 3)), "#dbe4f0")
        # page 1 may carry a cover of its own
        self.assertEqual(I._page_colour(_ir([[_fill(SHEET, "#123456")],
                                             [_fill(SHEET)], [_fill(SHEET)]])), "#dbe4f0")

    def test_a_sheet_page_one_alone_carries_is_a_cover(self):
        self.assertIsNone(I._page_colour(_ir([[_fill(SHEET)], [], []])))

    def test_one_later_sheet_alone_is_no_repetition(self):
        self.assertIsNone(I._page_colour(_ir([[], [_fill(SHEET)]])))
        self.assertEqual(I._page_colour(_ir([[_fill(SHEET)], [_fill(SHEET)]])), "#dbe4f0")

    def test_one_page_is_no_evidence(self):
        self.assertIsNone(I._page_colour(_ir([[_fill(SHEET)]])))

    def test_a_later_page_without_it_breaks_the_repetition(self):
        self.assertIsNone(I._page_colour(_ir([[_fill(SHEET)], [_fill(SHEET)], []])))

    def test_a_partial_fill_is_not_a_sheet(self):
        half = (0.0, 0.0, 612.0, 400.0)
        self.assertIsNone(I._page_colour(_ir([[_fill(half)]] * 3)))


def _span(text, x0, x1, base, size=10.0, font="Cambria", box=None):
    from exactdoc.model import Span
    top, bot = box if box else (base - 0.9 * size, base + 0.2 * size)
    return Span(text=text, font=font, size=size, color="#000000", bold=False,
                italic=False, mono=False, serif=True, superscript=False,
                bbox=(x0, top, x1, bot), origin=(x0, base))


def _text_ir(lines_per_page):
    from exactdoc.model import Line, TextBlock
    pages = []
    for i, lines in enumerate(lines_per_page):
        blocks = [TextBlock(lines=[ln], bbox=ln.bbox) for ln in lines]
        pages.append(PageIR(i + 1, 612.0, 792.0, blocks=blocks, drawings=[]))
    return DocIR(path="synthetic.pdf", pages=pages)


def _line(span, direction=(1.0, 0.0)):
    from exactdoc.model import Line
    return Line(spans=[span], bbox=span.bbox, dir=direction)


class TextFontBoxesUnit(unittest.TestCase):
    """`_true_text_boxes`: a text face with inflated boxes gets its glyphs'."""

    def _inflated(self, text, base, font="Cambria", direction=(1.0, 0.0)):
        # Word's Cambria: a box 5.6x the size, standing far above the baseline
        return _line(_span(text, 72.0, 300.0, base, font=font,
                           box=(base - 40.0, base + 16.0)), direction)

    def test_a_text_face_with_inflated_boxes_is_set_on_its_baselines(self):
        lines = [self._inflated("Carnegie Mellon University", 100.0 + 20 * k)
                 for k in range(6)]
        ir = _text_ir([lines])
        self.assertEqual(I._true_text_boxes(ir), 6)
        for k, ln in enumerate(lines):
            base = 100.0 + 20 * k
            self.assertAlmostEqual(ln.bbox[1], base - I.TEXT_ASCENT_EM * 10.0)
            self.assertAlmostEqual(ln.bbox[3], base + I.TEXT_DESCENT_EM * 10.0)
            self.assertEqual(ln.spans[0].origin[1], base)
        # each line's block follows its line
        self.assertEqual(ir.pages[0].blocks[0].bbox, lines[0].bbox)

    def test_a_maths_face_setting_one_glyph_at_a_time_keeps_its_boxes(self):
        lines = [self._inflated(g, 100.0 + 20 * k, font="CMEX10")
                 for k, g in enumerate("∑∫(){}")]
        before = [ln.bbox for ln in lines]
        self.assertEqual(I._true_text_boxes(_text_ir([lines])), 0)
        self.assertEqual([ln.bbox for ln in lines], before)

    def test_ordinary_boxes_are_left_alone(self):
        lines = [_line(_span("an ordinary line of text", 72.0, 300.0, 100.0 + 14 * k))
                 for k in range(8)]
        before = [ln.bbox for ln in lines]
        self.assertEqual(I._true_text_boxes(_text_ir([lines])), 0)
        self.assertEqual([ln.bbox for ln in lines], before)

    def test_a_rotated_line_keeps_its_box(self):
        lines = [self._inflated("Carnegie Mellon University", 100.0 + 20 * k)
                 for k in range(6)]
        rotated = self._inflated("Carnegie Mellon University", 400.0, direction=(0.0, 1.0))
        before = rotated.bbox
        self.assertEqual(I._true_text_boxes(_text_ir([lines + [rotated]])), 6)
        self.assertEqual(rotated.bbox, before)

    def test_switched_off_by_the_flag(self):
        lines = [self._inflated("Carnegie Mellon University", 100.0 + 20 * k)
                 for k in range(6)]
        before = [ln.bbox for ln in lines]
        I.infer(_text_ir([lines]), tab_layouts=False)
        self.assertEqual([ln.bbox for ln in lines], before)


class TallStripUnit(unittest.TestCase):
    """detect_hf: a running strip reaching past mid-page is not a header."""

    def _ir(self, strip_h):
        pages = []
        for n in range(4):
            line = _line(_span("Body text on page %d" % (n + 1), 72.0, 300.0, 700.0,
                               font="Times-Roman"))
            from exactdoc.model import TextBlock
            pages.append(PageIR(n + 1, 612.0, 792.0,
                                blocks=[TextBlock(lines=[line], bbox=line.bbox)],
                                drawings=[_fill((0.0, 0.0, 612.0, strip_h), "#1f487c")]))
        return DocIR(path="synthetic.pdf", pages=pages)

    def test_a_strip_past_mid_page_is_no_header(self):
        tok = I._TAB_LAYOUTS.set(True)
        try:
            with self.assertLogs("exactdoc.infer", "WARNING"):
                res = I.detect_hf(self._ir(500.0))
        finally:
            I._TAB_LAYOUTS.reset(tok)
        self.assertIsNone(res["band_def"])

    def test_a_real_header_strip_stays_a_strip(self):
        tok = I._TAB_LAYOUTS.set(True)
        try:
            res = I.detect_hf(self._ir(60.0))
        finally:
            I._TAB_LAYOUTS.reset(tok)
        self.assertIsNotNone(res["band_def"])

    def test_with_the_flag_off_the_tall_strip_is_read_as_before(self):
        tok = I._TAB_LAYOUTS.set(False)
        try:
            res = I.detect_hf(self._ir(500.0))
        finally:
            I._TAB_LAYOUTS.reset(tok)
        self.assertIsNotNone(res["band_def"])


class SectionBackstop(unittest.TestCase):
    def test_margins_that_leave_no_body_are_clamped_and_counted(self):
        from exactdoc.layout import DocLayout
        from docx import Document
        lay = DocLayout()
        lay.margin_t, lay.margin_b = 792.0, 72.0
        before = docxout.SECTION_CLAMPS[0]
        sec = Document().sections[0]
        with self.assertLogs("exactdoc.docxout", "WARNING"):
            docxout._config_section(sec, lay)
        self.assertEqual(docxout.SECTION_CLAMPS[0], before + 1)
        self.assertLessEqual(sec.top_margin.pt + sec.bottom_margin.pt,
                             lay.page_h - docxout.MIN_BODY_PT + 0.01)

    def test_a_bottom_margin_that_alone_leaves_no_body_is_clamped_too(self):
        # WP64 (S4): with margin_b >= page_h - one line, clamping the top
        # alone still left the section no body, and Word refuses the file.
        from exactdoc.layout import DocLayout
        from docx import Document
        for top, bottom in ((72.0, 790.0), (0.0, 792.0), (400.0, 785.0)):
            with self.subTest(top=top, bottom=bottom):
                lay = DocLayout()
                lay.margin_t, lay.margin_b = top, bottom
                sec = Document().sections[0]
                with self.assertLogs("exactdoc.docxout", "WARNING"):
                    docxout._config_section(sec, lay)
                body = lay.page_h - sec.top_margin.pt - sec.bottom_margin.pt
                self.assertGreaterEqual(body, docxout.MIN_BODY_PT - 0.01)
                self.assertGreaterEqual(sec.top_margin.pt, 0.0)

    def test_sane_margins_are_written_as_they_are(self):
        from exactdoc.layout import DocLayout
        from docx import Document
        lay = DocLayout()
        lay.margin_t, lay.margin_b = 54.0, 61.5
        before = docxout.SECTION_CLAMPS[0]
        sec = Document().sections[0]
        docxout._config_section(sec, lay)
        self.assertEqual(docxout.SECTION_CLAMPS[0], before)
        self.assertAlmostEqual(sec.top_margin.pt, 54.0, places=2)
        self.assertAlmostEqual(sec.bottom_margin.pt, 61.5, places=2)

    def test_the_gdocs_switch_is_reset_after_every_write(self):
        from exactdoc.layout import DocLayout
        import tempfile
        with tempfile.TemporaryDirectory() as d:
            docxout.write_docx(DocLayout(), os.path.join(d, "g.docx"),
                               output_profile="gdocs")
        self.assertTrue(docxout._CLAMP_SECTIONS[0])


# --------------------------------------------------------- end to end
@unittest.skipUnless(F.available(), "reportlab is not installed")
class PageColourEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._dir = tempfile.TemporaryDirectory()
        d = cls._dir.name
        mk = lambda name, **kw: F.page_colour_document(os.path.join(d, name + ".pdf"), **kw)
        cls.paths = {
            "all": mk("all"),
            "all_but_cover": mk("all_but_cover", cover="#16497d"),
            "cover_only": mk("cover_only", colour=None, cover="#16497d", border=False),
            "single": mk("single", pages=1, border=False),
            "strip": mk("strip", strip="#1f487c", border=False),
            "dark": mk("dark", colour="#1f2a44", text_colour=(1, 1, 1), border=False),
        }

    @classmethod
    def tearDownClass(cls):
        cls._dir.cleanup()

    def out(self, name, lay):
        return _docx(lay, os.path.join(self._dir.name, name + ".docx"))

    def assert_sane_sections(self, xml):
        for h, mt, mb in _sections(xml):
            self.assertGreater(h - mt - mb, 0)
            self.assertLess(mt, h / 2)
            self.assertLess(mb, h / 2)

    def test_R1_page_colour_on_every_page_is_the_background(self):
        lay, old = _infer(self.paths["all"]), _infer(self.paths["all"], tab_layouts=False)
        # the defect, reproduced with WP59 off: a band/strip swallows the page
        self.assertTrue(old.cover_band is not None or old.margin_t > 396)
        self.assertEqual(lay.page_colour, "#dbe4f0")
        self.assertIsNone(lay.cover_band)
        self.assertLess(lay.margin_t, 200)
        for t, *_ in F.RESUME_TEXT:
            self.assertIn(t, _body_text(lay))
        xml = self.out("all", lay)
        self.assertIn('<w:background w:color="DBE4F0"/>', xml["word/document.xml"])
        self.assertIn("<w:displayBackgroundShape/>", xml["word/settings.xml"])
        self.assert_sane_sections(xml["word/document.xml"])
        # the page border is neither a box nor a rule
        els = [e for pl in lay.pages for ch in pl.chunks for e in ch.elements]
        self.assertFalse([e for e in els if isinstance(e, TableEl)])
        self.assertFalse([e for e in els if type(e).__name__ == "RuleEl"])

    def test_R1_page_colour_behind_a_cover(self):
        lay = _infer(self.paths["all_but_cover"])
        self.assertEqual(lay.page_colour, "#dbe4f0")
        self.assert_sane_sections(self.out("all_but_cover", lay)["word/document.xml"])

    def test_R1_controls_are_unchanged(self):
        for name in ("cover_only", "single"):
            with self.subTest(name):
                lay = _infer(self.paths[name])
                old = _infer(self.paths[name], tab_layouts=False)
                self.assertIsNone(lay.page_colour)
                self.assertEqual(_shape(lay), _shape(old))

    def test_R1_a_header_strip_under_a_page_colour_stays_a_strip(self):
        lay = _infer(self.paths["strip"])
        self.assertEqual(lay.page_colour, "#dbe4f0")
        self.assertIsNotNone(lay.header_default)
        self.assertLess(lay.margin_t, 120)
        self.assert_sane_sections(self.out("strip", lay)["word/document.xml"])

    def test_R1_a_dark_page_colour_keeps_its_white_text_on_it(self):
        lay = _infer(self.paths["dark"])
        self.assertEqual(lay.page_colour, "#1f2a44")
        xml = self.out("dark", lay)
        self.assertIn('<w:background w:color="1F2A44"/>', xml["word/document.xml"])
        self.assertIn('w:color w:val="FFFFFF"', xml["word/document.xml"])

    def test_R1_a_table_drawn_to_the_page_edge_is_no_page_border(self):
        path = F.full_page_table(os.path.join(self._dir.name, "table.pdf"))
        self.assertEqual(I._strip_page_border(_parse(path)), 0)
        lay, old = _infer(path), _infer(path, tab_layouts=False)
        self.assertEqual(_shape(lay), _shape(old))
        self.assertTrue([e for pl in lay.pages for ch in pl.chunks for e in ch.elements
                         if isinstance(e, TableEl)])

    def test_R1_flag_off_writes_no_background(self):
        lay = _infer(self.paths["all"], tab_layouts=False)
        self.assertIsNone(lay.page_colour)


@unittest.skipUnless(F.available(), "reportlab is not installed")
class RowsEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import wp49_fixtures as F49
        cls._dir = tempfile.TemporaryDirectory()
        d = cls._dir.name
        cls.paths = {n: getattr(F, n)(os.path.join(d, n + ".pdf"))
                     for n in ("two_sided_list_resume", "left_tab_resume",
                               "left_number_contents", "two_column_short_lines")}
        cls.paths["right_aligned_sidebar"] = F49.right_aligned_sidebar(
            os.path.join(d, "sidebar.pdf"))
        cls.paths["right_tab_resume"] = F49.right_tab_resume(os.path.join(d, "rt.pdf"))

    @classmethod
    def tearDownClass(cls):
        cls._dir.cleanup()

    def paras(self, name, **kw):
        lay = _infer(self.paths[name], **kw)
        return [e for pl in lay.pages for ch in pl.chunks for e in ch.elements
                if isinstance(e, Para)]

    def test_R2_a_two_sided_list_keeps_its_rows(self):
        texts = [p.text for p in self.paras("two_sided_list_resume")]
        for left, right in F.R2_LIST:
            self.assertIn(left + "	" + right, texts)

    def test_R3a_left_tab_fields_are_rows_with_a_left_stop(self):
        ps = self.paras("left_tab_resume")
        rows = [p for p in ps if "	" in p.text and not p.text.startswith("•")]
        for label, field in F.R3_ROWS:
            self.assertIn(label + "	" + field, [p.text for p in rows])
        for p in rows:
            self.assertEqual(len(p.tab_stops), 1)
            self.assertEqual(p.tab_stops[0][1], "left")
            self.assertAlmostEqual(p.tab_stops[0][0], 430.0 - 72.0, delta=2.0)
        # with WP59 off, no row has a stop at the tab
        old = self.paras("left_tab_resume", tab_layouts=False)
        self.assertFalse([p for p in old if p.tab_stops and p.tab_stops[0][1] == "left"
                          and "	" in p.text and not p.text.startswith("•")])

    def test_R3a_controls_are_unchanged(self):
        for name in ("left_number_contents", "two_column_short_lines",
                     "right_aligned_sidebar", "right_tab_resume"):
            with self.subTest(name):
                lay = _infer(self.paths[name])
                if name == "right_tab_resume":
                    # WP49's rows, unchanged by WP59 (a WP59 switch off)
                    off = dict(I.WP59_RULES)
                    try:
                        for k in I.WP59_RULES:
                            I.WP59_RULES[k] = False
                        base = _infer(self.paths[name])
                    finally:
                        I.WP59_RULES.update(off)
                    self.assertEqual(_shape(lay), _shape(base))
                    continue
                old = _infer(self.paths[name], tab_layouts=False)
                self.assertEqual(_shape(lay), _shape(old))


if __name__ == "__main__":
    unittest.main()
